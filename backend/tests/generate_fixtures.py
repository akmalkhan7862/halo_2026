"""
Fixture Generator Script
Generates clean_resume.pdf, clean_resume.docx, complex_resume.pdf, and sample_jds.json.
"""
import os
import json
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

FIXTURES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tests", "fixtures")
os.makedirs(FIXTURES_DIR, exist_ok=True)


def generate_clean_pdf():
    pdf_path = os.path.join(FIXTURES_DIR, "clean_resume.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("<b>Alex Morgan</b>", styles["Title"]))
    story.append(Paragraph("alex.morgan@example.com | (555) 234-5678 | San Francisco, CA | github.com/alexmorgan | linkedin.com/in/alexmorgan", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", styles["Heading2"]))
    story.append(Paragraph("Results-driven Software Engineer with 4 years of hands-on experience building high-throughput backend APIs, microservices, and distributed data pipelines.", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>TECHNICAL SKILLS</b>", styles["Heading2"]))
    story.append(Paragraph("Languages: Python, TypeScript, SQL", styles["Normal"]))
    story.append(Paragraph("Frameworks: FastAPI, Flask, Pydantic, SQLAlchemy", styles["Normal"]))
    story.append(Paragraph("Databases: PostgreSQL, Redis, MySQL", styles["Normal"]))
    story.append(Paragraph("Cloud & DevOps: Docker, Git, CI/CD, AWS", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>WORK EXPERIENCE</b>", styles["Heading2"]))
    story.append(Paragraph("<b>Software Engineer | Apex Cloud Systems</b> (Jan 2022 - Present)", styles["Normal"]))
    story.append(Paragraph("* Developed REST APIs using FastAPI and PostgreSQL supporting over 50,000 daily active users.", styles["Normal"]))
    story.append(Paragraph("* Optimized database indexing and async queries, reducing p95 API response times by 38%.", styles["Normal"]))
    story.append(Paragraph("* Architected automated CI/CD deployment pipelines using Docker and GitHub Actions.", styles["Normal"]))
    story.append(Spacer(1, 8))

    story.append(Paragraph("<b>Junior Developer | NovaTech Labs</b> (Jun 2020 - Dec 2021)", styles["Normal"]))
    story.append(Paragraph("* Maintained backend Flask microservices and integrated third-party payment gateways.", styles["Normal"]))
    story.append(Paragraph("* Implemented unit and integration test suites achieving 85% code coverage.", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>KEY PROJECTS</b>", styles["Heading2"]))
    story.append(Paragraph("<b>Distributed Task Streamer</b>", styles["Normal"]))
    story.append(Paragraph("Engineered a resilient event streaming worker using Python, Redis, and Docker. Handled 10,000 tasks/sec with zero job loss.", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>EDUCATION</b>", styles["Heading2"]))
    story.append(Paragraph("Bachelor of Science in Computer Science | University of California, Berkeley (2020) | GPA: 3.8", styles["Normal"]))

    doc.build(story)
    print(f"Generated {pdf_path}")


def generate_clean_docx():
    docx_path = os.path.join(FIXTURES_DIR, "clean_resume.docx")
    doc = docx.Document()

    doc.add_heading("Taylor Reed", level=0)
    doc.add_paragraph("taylor.reed@example.com | (555) 987-6543 | github.com/taylorreed")

    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Full Stack Developer with expertise in React, TypeScript, Python, and cloud infrastructure.")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph("React, TypeScript, JavaScript, Python, FastAPI, Docker, PostgreSQL, Git")

    doc.add_heading("Experience", level=1)
    p1 = doc.add_paragraph("Full Stack Engineer at SkyLine Inc. (Mar 2021 - Present)")
    doc.add_paragraph("* Built interactive frontend dashboards using React and TypeScript.")
    doc.add_paragraph("* Integrated backend microservices built in FastAPI and PostgreSQL.")

    doc.add_heading("Projects", level=1)
    doc.add_paragraph("Real-Time Collaboration Board: Web application built with React, Node.js, and Docker.")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Technology in Information Technology (2021)")

    doc.save(docx_path)
    print(f"Generated {docx_path}")


def generate_complex_pdf():
    pdf_path = os.path.join(FIXTURES_DIR, "complex_resume.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("<b>Jordan Lee, Ph.D.</b>", styles["Title"]))
    story.append(Paragraph("jordan.lee@example.com | San Jose, CA", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>CORE COMPETENCIES & TECHNOLOGIES</b>", styles["Heading2"]))
    story.append(Paragraph("Distributed Systems, System Design, Python, C++, Kubernetes, Docker, PostgreSQL, Redis, Microservices", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>PROFESSIONAL HISTORY</b>", styles["Heading2"]))
    story.append(Paragraph("Staff Infrastructure Architect | GlobalScale Corp (2019 - Present)", styles["Normal"]))
    story.append(Paragraph("* Spearheaded redesign of core transaction ledger handling $10M+ in daily transaction volume.", styles["Normal"]))
    story.append(Paragraph("* Mentored 12 senior engineers in distributed consensus and system resiliency patterns.", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>ACADEMIC BACKGROUND</b>", styles["Heading2"]))
    story.append(Paragraph("Ph.D. in Computer Engineering | Stanford University (2018)", styles["Normal"]))

    doc.build(story)
    print(f"Generated {pdf_path}")


def generate_sample_jds():
    jds_path = os.path.join(FIXTURES_DIR, "sample_jds.json")
    data = [
        {
            "id": "jd-backend-standard",
            "title": "Senior Backend Developer",
            "company": "CloudWave Technologies",
            "raw_text": """
CloudWave Technologies is hiring a Senior Backend Developer.

About the Role:
We are looking for an experienced Senior Backend Developer to join our core platform engineering team.
You will architect, build, and deploy mission-critical microservices and database pipelines.

Requirements:
* 4+ years of professional software engineering experience.
* Proficiency in Python and FastAPI (or modern Python web framework).
* Strong relational database expertise with PostgreSQL and complex SQL queries.
* Hands-on experience with Docker, containerization, and automated CI/CD workflows.
* Familiarity with version control using Git.

Preferred Qualifications:
* Experience with Kubernetes container orchestration.
* Knowledge of Amazon Web Services (AWS) cloud architecture.
* Understanding of System Design for scalable distributed services.
* Experience with caching layers such as Redis.

Responsibilities:
* Design scalable REST APIs and high-throughput background processing services.
* Collaborate closely with product managers and frontend teams.
* Conduct architecture reviews and mentor junior developers.
"""
        },
        {
            "id": "jd-short-minimal",
            "title": "Python Developer",
            "company": "FastStartup",
            "raw_text": "Looking for Python developer to write backend code in FastAPI and PostgreSQL with Docker."
        },
        {
            "id": "jd-no-explicit-skills",
            "title": "Technical Problem Solver",
            "company": "InnovateCo",
            "raw_text": "Join our collaborative team to solve tough customer problems, write reliable software, and deliver results."
        }
    ]

    with open(jds_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    print(f"Generated {jds_path}")


if __name__ == "__main__":
    generate_clean_pdf()
    generate_clean_docx()
    generate_complex_pdf()
    generate_sample_jds()
